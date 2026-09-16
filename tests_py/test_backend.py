from __future__ import annotations

from pathlib import Path
from io import BytesIO
import hashlib
import json

import pytest

from backend.app import create_app
from backend.agent import AgentOrchestrator, PlanValidationError, parse_action_plan
from backend.config import ROOT_DIR
from backend.config import default_config
from backend.inference import InferenceAdapter
from backend.model_download import ModelDownloadError, ModelReleaseDownloader
from backend.parsers import ParseError, parse_upload
from backend.rag import FaissRagService


@pytest.fixture()
def app(tmp_path: Path):
    application = create_app(
        {
            "ROOT_DIR": ROOT_DIR,
            "DATA_DIR": tmp_path / "data",
            "INDEX_DIR": tmp_path / "indexes",
            "MODELS_DIR": tmp_path / "models",
            "DATABASE_PATH": tmp_path / "data" / "test.db",
            "EMBEDDING_MODEL_PATH": "",
            "LLAMA_CPP_ENABLED": False,
            "OCR_ENABLED": False,
            "TESTING": True,
        }
    )
    yield application


@pytest.fixture()
def client(app):
    return app.test_client()


def test_health_and_faiss_index_are_ready(client, app, tmp_path: Path):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.get_json()["status"] == "ok"
    assert (tmp_path / "indexes" / "knowledge.faiss").exists()
    runtime = client.get("/api/v1/runtime/status").get_json()
    assert runtime["rag"]["chunkCount"] > 0
    assert runtime["rag"]["embedding"]["mode"] == "offline-fallback"
    assert runtime["ingestion"]["formats"] == ["txt", "md", "pdf", "docx", "xlsx", "csv", "pptx", "html"]
    assert runtime["ingestion"]["ocr"]["enabled"] is False
    assert runtime["modelDownload"]["repository"] == "liangyanlun/edge-office"


class _DownloadResponse:
    def __init__(self, body: bytes):
        self._body = body
        self._offset = 0
        self.headers = {"Content-Length": str(len(body))}

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def read(self, size: int = -1) -> bytes:
        if size < 0:
            size = len(self._body) - self._offset
        result = self._body[self._offset:self._offset + size]
        self._offset += len(result)
        return result


def test_model_release_downloader_installs_only_verified_release_assets(tmp_path: Path):
    models_dir = tmp_path / "models"
    model_path = models_dir / "qwen" / "Qwen3.5-0.8B-office.Q8_0.gguf"
    manifest_path = model_path.with_name("release.manifest.json")
    checksum_path = model_path.with_name(f"{model_path.stem}.SHA256SUMS.txt")
    model_bytes = b"verified-gguf"
    manifest_bytes = json.dumps({
        "schema": "edge_office_model_release_v1",
        "model_file": model_path.name,
        "sha256": hashlib.sha256(model_bytes).hexdigest(),
        "quantization": "Q8_0",
        "context_length": 2048,
        "acceptance_passed": True,
    }).encode("utf-8")
    base = "https://github.com/liangyanlun/edge-office/releases/download/v0.1.0"
    release_url = "https://api.github.com/repos/liangyanlun/edge-office/releases/latest"
    release_bytes = json.dumps({"tag_name": "v0.1.0", "assets": [
        {"name": model_path.name, "size": len(model_bytes), "browser_download_url": f"{base}/{model_path.name}"},
        {"name": manifest_path.name, "size": len(manifest_bytes), "browser_download_url": f"{base}/{manifest_path.name}"},
    ]}).encode("utf-8")
    payloads = {release_url: release_bytes, f"{base}/{model_path.name}": model_bytes, f"{base}/{manifest_path.name}": manifest_bytes}

    def opener(request, timeout):
        assert timeout == 30
        return _DownloadResponse(payloads[request.full_url])

    downloader = ModelReleaseDownloader(
        enabled=True, repository="liangyanlun/edge-office", release_tag="", models_dir=models_dir,
        model_path=model_path, manifest_path=manifest_path, checksum_path=checksum_path, n_ctx=2048, release_required=True, opener=opener,
    )
    result = downloader.install_sync()
    assert result["state"] == "completed"
    assert result["verification"] == "release-manifest-sha256"
    assert model_path.read_bytes() == model_bytes
    assert json.loads(manifest_path.read_text(encoding="utf-8"))["sha256"] == hashlib.sha256(model_bytes).hexdigest()


def test_model_release_downloader_rejects_untrusted_asset_url(tmp_path: Path):
    models_dir = tmp_path / "models"
    model_path = models_dir / "qwen" / "Qwen3.5-0.8B-office.Q8_0.gguf"
    manifest_path = model_path.with_name("release.manifest.json")
    checksum_path = model_path.with_name(f"{model_path.stem}.SHA256SUMS.txt")
    release_url = "https://api.github.com/repos/liangyanlun/edge-office/releases/latest"
    release_bytes = json.dumps({"tag_name": "v0.1.0", "assets": [
        {"name": model_path.name, "size": 12, "browser_download_url": "https://example.invalid/model.gguf"},
    ]}).encode("utf-8")

    def opener(request, timeout):
        return _DownloadResponse(release_bytes if request.full_url == release_url else b"")

    downloader = ModelReleaseDownloader(
        enabled=True, repository="liangyanlun/edge-office", release_tag="", models_dir=models_dir,
        model_path=model_path, manifest_path=manifest_path, checksum_path=checksum_path, n_ctx=2048, release_required=False, opener=opener,
    )
    with pytest.raises(ModelDownloadError, match="不受信任"):
        downloader.install_sync()
    assert not model_path.exists()


def test_model_release_downloader_accepts_matching_sha256sums_when_manifest_is_absent(tmp_path: Path):
    models_dir = tmp_path / "models"
    model_path = models_dir / "qwen" / "Qwen3.5-0.8B-office.Q8_0.gguf"
    manifest_path = model_path.with_name("release.manifest.json")
    checksum_path = model_path.with_name(f"{model_path.stem}.SHA256SUMS.txt")
    model_bytes = b"checksum-verified-gguf"
    checksum_bytes = f"{hashlib.sha256(model_bytes).hexdigest()}  {model_path.name}\n".encode("utf-8")
    base = "https://github.com/liangyanlun/edge-office/releases/download/v0.1.0"
    release_url = "https://api.github.com/repos/liangyanlun/edge-office/releases/latest"
    release_bytes = json.dumps({"tag_name": "v0.1.0", "assets": [
        {"name": model_path.name, "size": len(model_bytes), "browser_download_url": f"{base}/{model_path.name}"},
        {"name": checksum_path.name, "size": len(checksum_bytes), "browser_download_url": f"{base}/{checksum_path.name}"},
    ]}).encode("utf-8")
    payloads = {release_url: release_bytes, f"{base}/{model_path.name}": model_bytes, f"{base}/{checksum_path.name}": checksum_bytes}

    def opener(request, timeout):
        return _DownloadResponse(payloads[request.full_url])

    downloader = ModelReleaseDownloader(
        enabled=True, repository="liangyanlun/edge-office", release_tag="", models_dir=models_dir,
        model_path=model_path, manifest_path=manifest_path, checksum_path=checksum_path,
        n_ctx=2048, release_required=True, opener=opener,
    )
    result = downloader.install_sync()
    assert result["verification"] == "sha256sums"
    assert checksum_path.read_bytes() == checksum_bytes

    adapter = InferenceAdapter(
        None, model_path=model_path, release_manifest=manifest_path, release_checksum=checksum_path,
        n_ctx=2048, release_required=True,
    )  # type: ignore[arg-type]
    assert adapter._verify_release_manifest()["verification"] == "sha256sums"


def test_pinned_release_download_does_not_call_the_rate_limited_github_api(tmp_path: Path):
    models_dir = tmp_path / "models"
    model_path = models_dir / "qwen" / "Qwen3.5-0.8B-office.Q8_0.gguf"
    manifest_path = model_path.with_name("release.manifest.json")
    checksum_path = model_path.with_name(f"{model_path.stem}.SHA256SUMS.txt")
    model_bytes = b"pinned-release-gguf"
    checksum_bytes = f"{hashlib.sha256(model_bytes).hexdigest()}  {model_path.name}\n".encode("utf-8")
    base = "https://github.com/liangyanlun/edge-office/releases/download/v0.1.0"
    payloads = {f"{base}/{model_path.name}": model_bytes, f"{base}/{checksum_path.name}": checksum_bytes}

    def opener(request, timeout):
        assert "/api.github.com/" not in request.full_url
        return _DownloadResponse(payloads[request.full_url])

    downloader = ModelReleaseDownloader(
        enabled=True, repository="liangyanlun/edge-office", release_tag="v0.1.0", models_dir=models_dir,
        model_path=model_path, manifest_path=manifest_path, checksum_path=checksum_path,
        n_ctx=2048, release_required=True, opener=opener,
    )
    assert downloader.install_sync()["verification"] == "sha256sums"


def test_model_download_endpoint_rejects_client_supplied_urls(client):
    response = client.post("/api/v1/model/download", json={"url": "https://example.invalid/model.gguf"})
    assert response.status_code == 400
    assert response.get_json()["error"]["code"] == "UNTRUSTED_MODEL_DOWNLOAD_PAYLOAD"


def test_faiss_index_can_be_written_under_a_chinese_path(tmp_path: Path):
    unicode_root = tmp_path / "大创索引"
    application = create_app(
        {
            "ROOT_DIR": ROOT_DIR,
            "DATA_DIR": unicode_root / "数据",
            "INDEX_DIR": unicode_root / "索引",
            "MODELS_DIR": unicode_root / "模型",
            "DATABASE_PATH": unicode_root / "数据" / "test.db",
            "EMBEDDING_MODEL_PATH": "",
            "LLAMA_CPP_ENABLED": False,
            "TESTING": True,
        }
    )
    assert application.test_client().get("/api/v1/health").status_code == 200
    assert (unicode_root / "索引" / "knowledge.faiss").exists()


def test_config_honors_explicit_runtime_root(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    runtime_root = tmp_path / "edge-office-user-data"
    monkeypatch.setenv("EDGE_OFFICE_HOME", str(runtime_root))
    config = default_config()
    assert config["RUNTIME_DIR"] == runtime_root
    assert config["DATABASE_PATH"] == runtime_root / "artifacts" / "data" / "edge_office.db"
    assert config["MODELS_DIR"] == runtime_root / "artifacts" / "models"


def test_preferences_are_local_persistent_and_validate_input(client):
    initial = client.get("/api/v1/preferences")
    assert initial.status_code == 200
    assert initial.get_json()["item"]["onboardingComplete"] is False

    saved = client.put("/api/v1/preferences", json={
        "displayName": "言伦", "role": "科研学习", "writingTone": "concise",
        "ragEnabled": False, "onboardingComplete": True, "pinnedTools": ["summary"], "lastView": "tools",
    })
    assert saved.status_code == 200
    assert saved.get_json()["item"]["displayName"] == "言伦"
    assert saved.get_json()["item"]["pinnedTools"] == ["summary"]
    assert client.get("/api/v1/preferences").get_json()["item"]["lastView"] == "tools"

    invalid = client.put("/api/v1/preferences", json={"writingTone": "anything"})
    assert invalid.status_code == 400
    assert invalid.get_json()["error"]["code"] == "INVALID_PREFERENCES"


def test_document_import_builds_faiss_and_returns_cited_sse(client):
    imported = client.post(
        "/api/v1/documents",
        json={"name": "中期检查通知.md", "content": "中期检查材料需要在周五前提交，并附上实验记录和运行截图。"},
    )
    assert imported.status_code == 201
    assert imported.get_json()["index"]["chunks"] >= 5

    stream = client.post("/api/v1/chat/stream", json={"message": "中期检查材料什么时候提交？", "ragEnabled": True})
    body = stream.data.decode("utf-8")
    assert stream.status_code == 200
    assert "event: plan" in body
    assert "event: tool_start" in body
    assert "event: citations" in body
    assert "中期检查通知.md" in body
    assert "周五前提交" in body
    assert "event: done" in body


def test_agent_lists_documents_and_persists_the_conversation(client):
    response = client.post("/api/v1/agent/runs", json={"message": "有哪些材料可以查看？", "ragEnabled": True})
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["run"]["action"] == "list_documents"
    assert "项目综述" in payload["answer"]
    conversation = client.get(f"/api/v1/conversations/{payload['conversationId']}").get_json()["item"]
    assert len(conversation["messages"]) == 2


def test_chat_can_hide_internal_task_prompt_from_conversation_history(client):
    response = client.post(
        "/api/v1/chat/stream",
        json={
            "message": "请从本地材料中定位与这个问题最相关的原文，并保留引用依据：项目的核心创新是什么？",
            "displayMessage": "项目的核心创新是什么？",
            "ragEnabled": True,
        },
    )
    assert response.status_code == 200
    body = response.data.decode("utf-8")
    conversation_id = json.loads(next(line[6:] for line in body.splitlines() if line.startswith("data: ")))["conversationId"]
    conversation = client.get(f"/api/v1/conversations/{conversation_id}").get_json()["item"]
    assert conversation["messages"][0]["content"] == "项目的核心创新是什么？"
    assert "请从本地材料中定位" not in conversation["messages"][0]["content"]


def test_rag_persists_provenance_and_atomic_generation(app, client, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    imported = client.post(
        "/api/v1/documents",
        json={"name": "保密事项.md", "content": "内部代号 ORION 只允许项目组成员查看。", "securityLevel": "internal"},
    )
    assert imported.status_code == 201
    item = imported.get_json()["item"]
    assert item["versionId"].startswith(f"{item['id']}@")
    assert len(item["fileSha256"]) == 64
    assert item["securityLevel"] == "internal"

    rag = app.extensions["rag"]
    first_generation = rag.status()["activeGeneration"]
    assert (tmp_path / "indexes" / "active.json").exists()
    assert (tmp_path / "indexes" / "generations" / first_generation / "manifest.json").exists()
    assert rag.search("ORION", top_k=1)[0]["chunkId"].startswith(item["versionId"])

    original_encode_passages = rag.embedder.encode_passages
    monkeypatch.setattr(rag.embedder, "encode_passages", lambda _: (_ for _ in ()).throw(RuntimeError("embedding failed")))
    with pytest.raises(RuntimeError, match="embedding failed"):
        rag.rebuild()
    monkeypatch.setattr(rag.embedder, "encode_passages", original_encode_passages)
    assert rag.status()["activeGeneration"] == first_generation
    assert rag.search("ORION", top_k=1)[0]["documentId"] == item["id"]


def test_rag_filters_before_scoring_rejects_low_confidence_and_audits(app, client):
    internal = client.post(
        "/api/v1/documents",
        json={"name": "内部计划.md", "content": "ORION-77 将于周五提交，只供内部查看。", "securityLevel": "internal"},
    ).get_json()["item"]
    public = client.post(
        "/api/v1/documents",
        json={"name": "公开说明.md", "content": "本项目面向日常办公场景。", "securityLevel": "public"},
    ).get_json()["item"]
    rag = app.extensions["rag"]

    denied = rag.search("ORION-77 什么时候提交？", top_k=2, allowed_security_levels={"public"}, request_id="denied-request")
    assert denied == []
    allowed = rag.search("ORION-77 什么时候提交？", top_k=2, document_ids=[internal["id"]], request_id="allowed-request")
    assert allowed and all(item["documentId"] == internal["id"] for item in allowed)
    assert all("chunkId" in item and "quoteSha256" in item and "quoteEnd" in item for item in allowed)
    assert rag.search("完全不存在的编号 ZXQ-9999", top_k=2) == []

    database = app.extensions["database"]
    with database.connect() as connection:
        denied_event = connection.execute("SELECT status FROM retrieval_events WHERE request_id = ?", ("denied-request",)).fetchone()
        allowed_event = connection.execute("SELECT status, evidence_json FROM retrieval_events WHERE request_id = ?", ("allowed-request",)).fetchone()
    assert denied_event["status"] == "insufficient_confidence"
    assert allowed_event["status"] == "ok"
    assert "ORION-77" not in allowed_event["evidence_json"]
    assert public["id"] != internal["id"]


def test_rag_injects_at_most_four_blocks_within_frozen_budget(app):
    rag = app.extensions["rag"]
    evidence = rag.search("项目 目标 RAG 实施路线 性能", top_k=20)
    assert len(evidence) <= 4
    total = sum(
        24 + rag._token_cost(item["fileName"]) + rag._token_cost(item["locator"])
        + max(rag._token_cost(item["quote"]), rag._token_cost(item["context"]))
        for item in evidence
    )
    assert total <= 900
    assert rag.status()["retrieval"]["candidateK"] == 12
    assert rag.status()["retrieval"]["rerankK"] == 5


def test_rag_conflicting_cross_document_deadlines_fail_closed(app):
    rag = app.extensions["rag"]
    evidence = [
        {"documentId": "a", "score": 0.91, "quote": "材料必须在周五提交。"},
        {"documentId": "b", "score": 0.88, "quote": "材料必须在周一提交。"},
    ]
    assert rag._evidence_conflicts("材料什么时候提交？", evidence) is True
    evidence[1]["documentId"] = "a"
    assert rag._evidence_conflicts("材料什么时候提交？", evidence) is False


def test_model_citation_validation_only_keeps_returned_evidence():
    evidence = [{"chunkId": "doc@v#chunk-1", "quote": "证据"}]
    answer = InferenceAdapter._ensure_citations("模型错误标注【9】。", evidence)
    assert "【9】" not in answer
    assert "【1】" in answer


def test_model_citation_validation_normalizes_square_brackets():
    evidence = [{"chunkId": "doc@v#chunk-1", "quote": "evidence"}]
    answer = InferenceAdapter._ensure_citations("Model cites [1] and invalid [9].", evidence)
    assert "【1】" in answer
    assert "[1]" not in answer
    assert "[9]" not in answer


def test_llama_metrics_use_reported_completion_tokens_and_real_inference_time(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    class FakeModel:
        def create_chat_completion(self, **_kwargs):
            return {
                "choices": [{"message": {"content": "测试回答"}}],
                "usage": {"prompt_tokens": 40, "completion_tokens": 20},
            }

    adapter = InferenceAdapter(None, model_path=tmp_path / "fake.gguf")  # type: ignore[arg-type]
    monkeypatch.setattr(adapter, "_get_model", lambda: FakeModel())
    ticks = iter([10.0, 12.0])
    monkeypatch.setattr("backend.inference.time.perf_counter", lambda: next(ticks))

    assert adapter._chat("system", "user") == "测试回答"
    assert adapter.last_generation_metrics() == {
        "modelGenerationMs": 2000,
        "promptTokens": 40,
        "completionTokens": 20,
        "tokensPerSecond": 10.0,
        "failed": False,
    }


def test_stream_metrics_do_not_estimate_token_speed_from_character_count(client):
    stream = client.post("/api/v1/chat/stream", json={"message": "项目的核心创新是什么？", "ragEnabled": True})
    body = stream.data.decode("utf-8")
    metrics_line = next(line for line in body.splitlines() if line.startswith("data: ") and '"responseLatencyMs"' in line)
    metrics = json.loads(metrics_line.removeprefix("data: "))
    assert metrics["responseLatencyMs"] >= 1
    assert metrics["retrievalMs"] >= 1
    assert metrics["tokensPerSecond"] is None
    assert metrics["completionTokens"] is None


def test_document_search_binds_the_original_user_question(app, monkeypatch: pytest.MonkeyPatch):
    agent = app.extensions["agent"]
    monkeypatch.setattr(
        agent.inference,
        "create_action_plan",
        lambda *_: '{"tool":"document_search","arguments":{"query":"search-local-materials"},"confirmed":false}',
    )
    question = "What is the core innovation?"
    plan = agent.plan(question, None, True, "trusted-query-test")
    assert plan["tool"] == "document_search"
    assert plan["arguments"]["query"] == question


def test_rag_drops_low_relative_score_evidence():
    reranked = [
        (0.51, {"stable_id": "a"}, 0.0, 0.0),
        (0.38, {"stable_id": "b"}, 0.0, 0.0),
        (0.316, {"stable_id": "c"}, 0.0, 0.0),
        (0.129, {"stable_id": "d"}, 0.0, 0.0),
    ]
    selected = FaissRagService._select_evidence_candidates(reranked, top_k=4)
    assert [item[0] for item in selected] == [0.51, 0.38, 0.316]


def test_release_manifest_binds_hash_quantization_context_and_acceptance(tmp_path: Path):
    model = tmp_path / "Qwen3.5-0.8B.Q4_K_M.gguf"
    model.write_bytes(b"safe-test-gguf")
    manifest = tmp_path / "release.manifest.json"
    manifest.write_text(json.dumps({
        "schema": "edge_office_model_release_v1",
        "model_file": model.name,
        "sha256": hashlib.sha256(model.read_bytes()).hexdigest(),
        "quantization": "Q4_K_M",
        "context_length": 2048,
        "acceptance_passed": True,
    }), encoding="utf-8")
    adapter = InferenceAdapter(None, model_path=model, release_manifest=manifest, n_ctx=2048)  # type: ignore[arg-type]
    assert adapter._verify_release_manifest()["quantization"] == "Q4_K_M"
    model.write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="哈希"):
        adapter._verify_release_manifest()


def test_action_plan_rejects_invented_confirmation_unknown_fields_and_path_escape():
    with pytest.raises(PlanValidationError, match="模型不能确认"):
        parse_action_plan('{"tool":"email_send","arguments":{"to":"a@example.com","subject":"s","body":"b"},"confirmed":true}')
    with pytest.raises(PlanValidationError, match="只能包含"):
        parse_action_plan('{"tool":"document_search","arguments":{"query":"x"},"confirmed":false,"admin":true}')
    with pytest.raises(PlanValidationError, match="允许范围"):
        parse_action_plan('{"tool":"file_overwrite","arguments":{"path":"..\\\\secret.txt","content":"x"},"confirmed":false}')


def test_high_risk_plan_requires_bound_single_use_confirmation_and_is_audited(client, app):
    created = client.post("/api/v1/plans", json={"message": "请发送邮件给 team@example.com，主题是中期检查，正文是材料已准备完成。", "ragEnabled": False})
    assert created.status_code == 201
    plan = created.get_json()["item"]
    assert plan["tool"] == "email_send"
    assert plan["status"] == "AWAITING_CONFIRMATION"
    assert "confirmationId" not in plan and "confirmationNonce" not in plan
    untrusted = client.post(f"/api/v1/plans/{plan['id']}/confirm", json={"confirmed": True})
    assert untrusted.status_code == 400
    assert untrusted.get_json()["error"]["code"] == "UNTRUSTED_CONFIRMATION_PAYLOAD"
    other_session = app.test_client().post(f"/api/v1/plans/{plan['id']}/confirm")
    assert other_session.status_code == 409
    assert other_session.get_json()["error"]["code"] == "CONFIRMATION_SESSION_MISMATCH"

    confirmed = client.post(
        f"/api/v1/plans/{plan['id']}/confirm",
    )
    assert confirmed.status_code == 200
    assert confirmed.get_json()["plan"]["status"] == "SUCCEEDED"
    assert "沙箱" in confirmed.get_json()["answer"]
    replay = client.post(
        f"/api/v1/plans/{plan['id']}/confirm",
    )
    assert replay.status_code == 409

    audit = client.get(f"/api/v1/audit/{plan['requestId']}").get_json()
    assert audit["plans"][0]["status"] == "SUCCEEDED"
    assert {event["type"] for event in audit["plans"][0]["events"]} >= {"PLANNED", "VALIDATED", "CONFIRMED", "EXECUTING", "SUCCEEDED"}
    assert len(audit["plans"][0]["inputSha256"]) == 64
    assert len(audit["plans"][0]["modelOutputSha256"]) == 64
    assert audit["plans"][0]["toolAttempts"][0]["result"]["resultCharacters"] > 0
    assert "summary" not in audit["plans"][0]["toolAttempts"][0]["result"]
    assert audit["plans"][0]["confirmations"][0]["status"] == "CONSUMED"


def test_expired_confirmation_fails_closed(client, app):
    plan = client.post("/api/v1/plans", json={"message": "请发送邮件给 team@example.com，主题是中期检查，正文是材料已准备完成。"}).get_json()["item"]
    database = app.extensions["database"]
    with database.connect() as connection:
        connection.execute("UPDATE agent_plans SET expires_at = '2000-01-01T00:00:00+00:00' WHERE id = ?", (plan["id"],))
    response = client.post(
        f"/api/v1/plans/{plan['id']}/confirm",
    )
    assert response.status_code == 409
    assert app.extensions["database"].get_agent_plan(plan["id"])["status"] == "EXPIRED"


def test_trusted_intent_gate_never_broadens_a_small_model_tool_choice():
    assert AgentOrchestrator.candidate_tools("发送邮件给 alice@example.invalid，主题是项目进度，正文是已完成")[0]["name"] == "email_send"
    assert AgentOrchestrator.candidate_tools("帮我发送一封邮件")[0]["name"] == "request_clarification"
    assert AgentOrchestrator.candidate_tools("删除待办 task-12")[0]["name"] == "task_delete"
    assert AgentOrchestrator.candidate_tools("引用项目目标原文")[0]["name"] == "document_quote"
    draft_schema = AgentOrchestrator.candidate_tools("帮我写一封邮件草稿")[0]["schema"]
    assert set(draft_schema) == {"to", "subject", "body", "tone"}


def test_upload_import_parses_csv_html_and_preserves_source_locators(client, app):
    csv_file = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO("事项,截止时间\n提交中期材料,周五\n".encode("utf-8")), "任务.csv")},
        content_type="multipart/form-data",
    )
    assert csv_file.status_code == 201
    item = csv_file.get_json()["item"]
    assert item["mimeType"] == "text/csv"
    assert item["parserName"] == "csv"
    chunks = app.extensions["database"].list_chunks()
    assert any(chunk["document_id"] == item["id"] and "第 2 行" in chunk["locator"] for chunk in chunks)

    html_file = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(b"<h1>\xe9\xa1\xb9\xe7\x9b\xae\xe8\xbf\x9b\xe5\xba\xa6</h1><script>ignore()</script><p>\xe5\x91\xa8\xe4\xba\x94\xe6\x8f\x90\xe4\xba\xa4</p>"), "进度.html")},
        content_type="multipart/form-data",
    )
    assert html_file.status_code == 201
    assert html_file.get_json()["item"]["parserName"] == "html.parser"


def test_docx_and_scanned_pdf_parse_statuses(client):
    from docx import Document
    from PIL import Image, ImageDraw
    from pypdf import PdfWriter

    document = Document()
    document.add_heading("验收安排", level=1)
    document.add_paragraph("验收材料应于周五前提交。")
    payload = BytesIO()
    document.save(payload)
    docx_file = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(payload.getvalue()), "验收.docx")}, content_type="multipart/form-data",
    )
    assert docx_file.status_code == 201
    assert docx_file.get_json()["item"]["parserName"] == "python-docx"

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    empty_pdf = BytesIO()
    writer.write(empty_pdf)
    scanned = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(empty_pdf.getvalue()), "扫描件.pdf")}, content_type="multipart/form-data",
    )
    assert scanned.status_code == 400
    assert scanned.get_json()["error"]["code"] == "NO_EXTRACTABLE_TEXT"

    scan_image = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(scan_image).text((40, 80), "SCANNED DEADLINE FRIDAY", fill="black")
    scan_payload = BytesIO()
    scan_image.save(scan_payload, format="PDF", resolution=150)
    scan_image.close()
    ocr_disabled = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(scan_payload.getvalue()), "文字扫描件.pdf")}, content_type="multipart/form-data",
    )
    assert ocr_disabled.status_code == 400
    assert ocr_disabled.get_json()["error"]["code"] == "OCR_REQUIRED"


def test_xlsx_import_preserves_sheet_headers_and_cell_ranges(client, app):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目进度"
    sheet.append(["事项", "负责人", "截止日期"])
    sheet.append(["中期检查", "张三", "2026-09-01"])
    payload = BytesIO()
    workbook.save(payload)
    imported = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(payload.getvalue()), "项目进度.xlsx")}, content_type="multipart/form-data",
    )
    assert imported.status_code == 201
    item = imported.get_json()["item"]
    assert item["parserName"] == "openpyxl"
    assert item["parserMetadata"] == {"dataRowCount": 1, "sheetCount": 1}
    assert any(
        chunk["document_id"] == item["id"] and "工作表 项目进度 · A2:C2" in chunk["locator"] and "负责人：张三" in chunk["content"]
        for chunk in app.extensions["database"].list_chunks()
    )


def test_pptx_import_preserves_slide_table_and_notes(client, app):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "系统验收"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "材料", "截止"
    table.cell(1, 0).text, table.cell(1, 1).text = "实验截图", "周五"
    slide.notes_slide.notes_text_frame.text = "答辩时展示离线运行过程"
    payload = BytesIO()
    presentation.save(payload)
    imported = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(payload.getvalue()), "验收安排.pptx")}, content_type="multipart/form-data",
    )
    assert imported.status_code == 201
    item = imported.get_json()["item"]
    assert item["parserName"] == "python-pptx"
    assert item["parserMetadata"] == {"slideCount": 1}
    assert any(
        chunk["document_id"] == item["id"] and "第 1 页" in chunk["locator"] and "材料：实验截图" in chunk["content"]
        for chunk in app.extensions["database"].list_chunks()
    )


def test_office_limits_fail_before_database_write(client, app):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["列"])
    sheet.append(["数据"])
    payload = BytesIO()
    workbook.save(payload)
    app.config["MAX_XLSX_ROWS_PER_SHEET"] = 1
    before = len(app.extensions["database"].list_documents())
    rejected = client.post(
        "/api/v1/documents/import",
        data={"file": (BytesIO(payload.getvalue()), "超限.xlsx")}, content_type="multipart/form-data",
    )
    assert rejected.status_code == 400
    assert rejected.get_json()["error"]["code"] == "XLSX_ROW_LIMIT"
    assert len(app.extensions["database"].list_documents()) == before
