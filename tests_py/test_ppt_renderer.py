import base64
import json
import os
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from backend.ppt import renderer


def fixture(path):
    deck = Presentation()
    deck.slide_width = Inches(10)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text = '中文标题：项目总结'
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    deck.save(path)


def test_fallback_has_chinese_font_and_original_aspect_ratio(tmp_path):
    path = tmp_path / '中文含 空格.pptx'
    fixture(path)
    result = renderer.render_presentation(path, tmp_path / 'slides', False)
    assert result['approximate'] and result['warning']
    assert result['revision'] == renderer.RENDER_REVISION
    assert Image.open(tmp_path / 'slides/slide-1.png').size == (1280, 960)
    if os.name == 'nt':
        font = renderer.preview_font(32)
        assert bytes(font.getmask('中')) != bytes(font.getmask('国'))


def test_worker_hides_console_and_paths_are_not_commands(tmp_path, monkeypatch):
    seen = {}
    monkeypatch.setattr(renderer.shutil, 'which', lambda _: 'powershell.exe')
    def run(args, **kwargs):
        seen.update(args=args, **kwargs)
        directory = Path(kwargs['env']['EDGE_PPT_OUTPUT']); directory.mkdir()
        Image.new('RGB',(1280,960),'white').save(directory/'slide-1.png')
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(renderer.subprocess, 'run', run)
    path = tmp_path / "中文 $foo ' spaced.pptx"
    assert renderer._render_with_powerpoint_worker(path, tmp_path / 'previews')
    script = base64.b64decode(seen['args'][-1]).decode('utf-16-le')
    assert str(path) not in script and '$env:EDGE_PPT_INPUT' in script
    assert seen['env']['EDGE_PPT_INPUT'] == str(path.resolve())
    assert '-WindowStyle' in seen['args'] and 'Hidden' in seen['args']
    if os.name == 'nt':
        assert seen['creationflags'] & renderer.subprocess.CREATE_NO_WINDOW
        assert seen['startupinfo'].wShowWindow == 0


@pytest.mark.skipif(os.name != 'nt' or not os.getenv('EDGE_OFFICE_TEST_OFFICE'), reason='Explicit local Office integration check only')
def test_real_powerpoint_export(tmp_path):
    path = tmp_path / "原版 中文 $空格.pptx"
    fixture(path)
    result = renderer.render_presentation(path, tmp_path / 'slides')
    assert result['mode'] == 'powerpoint-com-worker', result
    assert Image.open(tmp_path / 'slides/slide-1.png').size == (1280, 960)
