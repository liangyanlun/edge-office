import subprocess
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from backend.ppt.renderer import hidden_process_options


def test_mobile_ooxml_roundtrip_with_real_pptx(tmp_path):
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2))
    box.text='原文';box.text_frame.add_paragraph().text='第二段'
    asset=tmp_path/'image.png';Image.new('RGB',(120,80),'green').save(asset)
    for left in (1,4):slide.shapes.add_picture(str(asset),Inches(left),Inches(3),Inches(2),Inches(1))
    slide2=deck.slides.add_slide(deck.slide_layouts[6])
    slide2.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2)).text='非目标页面'
    original=tmp_path/'original.pptx';deck.save(original)
    output=tmp_path/'result.pptx'
    import shutil
    node=shutil.which('node')
    assert node
    result=subprocess.run([node,str(Path(__file__).resolve().parent.parent/'tests/run-mobile-ppt-fixture.mjs'),str(original),str(output),str(asset)],capture_output=True,timeout=30,**hidden_process_options())
    assert result.returncode==0,result.stderr.decode('utf-8',errors='replace')
    reopened=Presentation(output)
    assert [p.text for p in reopened.slides[0].shapes[0].text_frame.paragraphs]==['手机新标题','手机新标题']
    assert len([s for s in reopened.slides[0].shapes if s.shape_type==13])==1
    assert reopened.slides[1].shapes[0].text=='非目标页面'
    assert Presentation(original).slides[0].shapes[0].text.startswith('原文')
