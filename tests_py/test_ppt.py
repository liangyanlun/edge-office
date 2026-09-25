from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend.database import Database
from backend.inference import InferenceAdapter, LocalModelRegistry
from backend.ppt.planner import _model_billions
from backend.ppt.routes import create_ppt_blueprint
from backend.ppt.service import PptService


@pytest.fixture()
def client(tmp_path: Path):
    database = Database(tmp_path / "data" / "test.db")
    database.initialize()
    inference = InferenceAdapter(LocalModelRegistry(tmp_path / "models"), llama_cpp_enabled=False)
    service = PptService(database, tmp_path / "data", inference, 20 * 1024 * 1024, allow_external_renderer=False)
    app = Flask(__name__)
    app.config["TESTING"] = True
    app.register_blueprint(create_ppt_blueprint(service))
    return app.test_client()


def pptx_bytes(with_picture: bool = False) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1.2))
    box.text_frame.paragraphs[0].text = "原始标题"
    box.text_frame.add_paragraph().text = "原始副标题"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "不应变化"
    if with_picture:
        for color, left in (("red", 1), ("blue", 4)):
            image = BytesIO()
            Image.new("RGB", (80, 50), color).save(image, format="PNG")
            image.seek(0)
            slide.shapes.add_picture(image, Inches(left), Inches(3), Inches(2), Inches(1.25))
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def import_deck(client, data: bytes):
    response = client.post(
        "/api/v1/presentations/import",
        data={"file": (BytesIO(data), "demo.pptx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 201, response.get_json()
    return response.get_json()["item"]


def test_ppt_versioned_edit_confirm_and_download(client):
    original = pptx_bytes()
    imported = import_deck(client, original)
    presentation_id = imported["id"]
    version_id = imported["currentVersionId"]
    snapshot = client.get(f"/api/v1/presentations/{presentation_id}/slides/slide-1/snapshot").get_json()["item"]
    text = next(item for item in snapshot["elements"] if item["kind"] == "text")
    sequence = f"replace_paragraph({text['id']}, 0, '2026 年项目总结')"
    created = client.post(
        f"/api/v1/presentations/{presentation_id}/patches",
        json={"baseVersion": version_id, "slideId": "slide-1", "instruction": "替换标题", "apiSequence": sequence},
    )
    assert created.status_code == 201, created.get_json()
    patch = created.get_json()["item"]
    assert patch["actionPlan"]["arguments"] == {"patch_id": patch["id"]}
    simulated = client.post(f"/api/v1/ppt-patches/{patch['id']}/simulate")
    assert simulated.status_code == 200, simulated.get_json()
    assert simulated.get_json()["item"]["verification"]["nonTargetSlidesVerified"] == 1
    forbidden = client.post(f"/api/v1/ppt-patches/{patch['id']}/confirm", json={"operations": []})
    assert forbidden.status_code == 400
    confirmed = client.post(f"/api/v1/ppt-patches/{patch['id']}/confirm")
    assert confirmed.status_code == 200, confirmed.get_json()
    new_version = confirmed.get_json()["version"]["id"]
    downloaded = client.get(f"/api/v1/presentations/{presentation_id}/versions/{new_version}/download")
    assert downloaded.status_code == 200
    reopened = Presentation(BytesIO(downloaded.data))
    assert "2026 年项目总结" in reopened.slides[0].shapes[0].text
    detail = client.get(f"/api/v1/presentations/{presentation_id}").get_json()["item"]
    assert detail["originalSha256"] == imported["originalSha256"]
    restored = client.post(f"/api/v1/presentations/{presentation_id}/versions/{version_id}/restore")
    assert restored.status_code == 200
    assert restored.get_json()["item"]["id"] != version_id


def test_ppt_rejects_arbitrary_image_path(client):
    imported = import_deck(client, pptx_bytes(with_picture=True))
    snapshot = client.get(f"/api/v1/presentations/{imported['id']}/slides/slide-1/snapshot").get_json()["item"]
    picture = next(item for item in snapshot["elements"] if item["kind"] == "image")
    response = client.post(
        f"/api/v1/presentations/{imported['id']}/patches",
        json={
            "baseVersion": imported["currentVersionId"], "slideId": "slide-1", "instruction": "替换图片",
            "apiSequence": f"replace_image({picture['id']}, 'C:/secret.png')",
        },
    )
    assert response.status_code == 422
    assert response.get_json()["error"]["code"] == "PPT_ASSET_URI_REQUIRED"


def test_all_five_ppt_operations_replay_deterministically(client):
    imported = import_deck(client, pptx_bytes(with_picture=True))
    presentation_id = imported["id"]
    snapshot = client.get(f"/api/v1/presentations/{presentation_id}/slides/slide-1/snapshot").get_json()["item"]
    text = next(item for item in snapshot["elements"] if item["kind"] == "text")
    pictures = [item for item in snapshot["elements"] if item["kind"] == "image"]
    replacement = BytesIO()
    Image.new("RGB", (120, 80), "green").save(replacement, format="PNG")
    replacement.seek(0)
    asset_response = client.post(
        f"/api/v1/presentations/{presentation_id}/assets",
        data={"file": (replacement, "replacement.png")},
        content_type="multipart/form-data",
    )
    assert asset_response.status_code == 201
    asset_uri = asset_response.get_json()["item"]["uri"]
    sequence = "\n".join([
        f"replace_paragraph({text['id']}, 0, '新标题')",
        f"clone_paragraph({text['id']}, 0)",
        f"del_paragraph({text['id']}, 1)",
        f"replace_image({pictures[0]['id']}, '{asset_uri}')",
        f"del_image({pictures[1]['id']})",
    ])
    created = client.post(
        f"/api/v1/presentations/{presentation_id}/patches",
        json={"baseVersion": imported["currentVersionId"], "slideId": "slide-1", "instruction": "执行五类操作", "apiSequence": sequence},
    )
    assert created.status_code == 201, created.get_json()
    patch_id = created.get_json()["item"]["id"]
    simulated = client.post(f"/api/v1/ppt-patches/{patch_id}/simulate")
    assert simulated.status_code == 200, simulated.get_json()
    assert simulated.get_json()["item"]["verification"]["operationCount"] == 5
    confirmed = client.post(f"/api/v1/ppt-patches/{patch_id}/confirm")
    assert confirmed.status_code == 200, confirmed.get_json()
    version_id = confirmed.get_json()["version"]["id"]
    deck = Presentation(BytesIO(client.get(f"/api/v1/presentations/{presentation_id}/versions/{version_id}/download").data))
    texts = [paragraph.text for paragraph in deck.slides[0].shapes[0].text_frame.paragraphs]
    assert texts == ["新标题", "新标题"]
    assert len([shape for shape in deck.slides[0].shapes if shape.shape_type == 13]) == 1


def test_model_size_detection_for_ppt_warning():
    assert _model_billions("Qwen3.5-0.8B-office.Q8_0") == 0.8
    assert _model_billions("qwen3_5_4B_ppt") == 4.0


def test_stale_ppt_patch_cannot_be_confirmed(client):
    imported = import_deck(client, pptx_bytes())
    presentation_id = imported["id"]
    snapshot = client.get(f"/api/v1/presentations/{presentation_id}/slides/slide-1/snapshot").get_json()["item"]
    text = next(item for item in snapshot["elements"] if item["kind"] == "text")
    created = client.post(
        f"/api/v1/presentations/{presentation_id}/patches",
        json={"baseVersion": imported["currentVersionId"], "slideId": "slide-1", "instruction": "替换标题", "apiSequence": f"replace_paragraph({text['id']}, 0, '旧计划')"},
    ).get_json()["item"]
    assert client.post(f"/api/v1/ppt-patches/{created['id']}/simulate").status_code == 200
    assert client.post(f"/api/v1/presentations/{presentation_id}/versions/{imported['currentVersionId']}/restore").status_code == 200
    response = client.post(f"/api/v1/ppt-patches/{created['id']}/confirm")
    assert response.status_code == 409
    assert response.get_json()["error"]["code"] == "PPT_BASE_VERSION_STALE"
