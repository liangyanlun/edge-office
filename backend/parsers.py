from __future__ import annotations

import csv
import io
import json
import re
import threading
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from pptx import Presentation


class ParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ParsedDocument:
    content: str
    mime_type: str
    parser_name: str
    parser_version: str = "1"
    locator: str = "正文"
    extraction_quality: str = "native"
    metadata: dict[str, Any] | None = None


MIME_BY_EXTENSION = {
    ".txt": "text/plain", ".md": "text/markdown", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".csv": "text/csv", ".html": "text/html", ".htm": "text/html",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

_OCR_ENGINE: Any | None = None
_OCR_LOCK = threading.Lock()


class _HtmlText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.items: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._skip += 1
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.items.append(f"[[LOCATOR:标题 {tag[1]}]]")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip:
            self._skip -= 1
        elif tag in {"p", "div", "li", "tr", "br"}:
            self.items.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip and data.strip():
            self.items.append(data.strip())


def parse_upload(
    filename: str,
    data: bytes,
    max_characters: int,
    *,
    ocr_enabled: bool = True,
    max_pdf_pages: int = 200,
    max_ocr_pages: int = 30,
    max_sheets: int = 30,
    max_sheet_rows: int = 5_000,
    max_sheet_columns: int = 100,
    max_slides: int = 300,
) -> ParsedDocument:
    safe_name = Path(filename).name
    if not safe_name or safe_name != filename or len(safe_name) > 180:
        raise ParseError("INVALID_FILENAME", "文件名不合法")
    extension = Path(safe_name).suffix.lower()
    if extension not in MIME_BY_EXTENSION:
        raise ParseError("UNSUPPORTED_FILE_TYPE", "仅支持 TXT、Markdown、PDF、DOCX、XLSX、CSV、PPTX 和 HTML")
    if not data:
        raise ParseError("EMPTY_FILE", "文件为空")
    try:
        if extension in {".txt", ".md"}:
            result = ParsedDocument(_decode_text(data), MIME_BY_EXTENSION[extension], "plain-text", locator="文本内容")
        elif extension == ".pdf":
            result = _parse_pdf(data, ocr_enabled=ocr_enabled, max_pages=max_pdf_pages, max_ocr_pages=max_ocr_pages)
        elif extension == ".docx":
            _validate_ooxml_archive(data)
            result = _parse_docx(data)
        elif extension == ".xlsx":
            _validate_ooxml_archive(data)
            result = _parse_xlsx(data, max_sheets=max_sheets, max_rows=max_sheet_rows, max_columns=max_sheet_columns)
        elif extension == ".pptx":
            _validate_ooxml_archive(data)
            result = _parse_pptx(data, max_slides=max_slides)
        elif extension == ".csv":
            result = _parse_csv(data)
        else:
            result = _parse_html(data)
    except ParseError:
        raise
    except Exception as error:
        raise ParseError("PARSER_FAILED", "文件解析失败，未写入知识库") from error
    content = re.sub(r"\n{3,}", "\n\n", result.content).strip()
    if not content:
        raise ParseError("NO_EXTRACTABLE_TEXT", "文件未提取到可检索文本；扫描 PDF 需要 OCR 支持")
    if len(content) > max_characters:
        raise ParseError("DOCUMENT_TOO_LARGE", "解析后的文本超过当前限制")
    return ParsedDocument(
        content, result.mime_type, result.parser_name, result.parser_version, result.locator,
        result.extraction_quality, result.metadata,
    )


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("TEXT_ENCODING_UNSUPPORTED", "文本编码不受支持，请使用 UTF-8 或 GB18030")


def _validate_ooxml_archive(data: bytes, max_files: int = 5_000, max_expanded_bytes: int = 100 * 1024 * 1024) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > max_files:
                raise ParseError("OOXML_FILE_LIMIT", "Office 文件内部项目过多")
            expanded = 0
            for item in entries:
                path = Path(item.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts:
                    raise ParseError("OOXML_PATH_ESCAPE", "Office 文件包含不安全的内部路径")
                expanded += item.file_size
                if expanded > max_expanded_bytes:
                    raise ParseError("OOXML_EXPANDED_LIMIT", "Office 文件解压后的内容超过 100MB 限制")
                if item.file_size > 1024 * 1024 and item.compress_size and item.file_size / item.compress_size > 200:
                    raise ParseError("OOXML_SUSPICIOUS_COMPRESSION", "Office 文件压缩比例异常")
    except zipfile.BadZipFile as error:
        raise ParseError("INVALID_OOXML", "Office 文件结构损坏") from error


def _parse_pdf(data: bytes, *, ocr_enabled: bool, max_pages: int, max_ocr_pages: int) -> ParsedDocument:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("ENCRYPTED_PDF", "加密 PDF 不能导入")
    if len(reader.pages) > max_pages:
        raise ParseError("PDF_PAGE_LIMIT", f"PDF 超过 {max_pages} 页限制")
    units_by_page: dict[int, str] = {}
    ocr_pages: list[int] = []
    blank_pages: list[int] = []
    confidence_values: list[float] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            units_by_page[index] = f"[[LOCATOR:第 {index} 页]]\n{text}"
        else:
            try:
                has_renderable_content = page.get_contents() is not None
            except Exception:
                has_renderable_content = True
            (ocr_pages if has_renderable_content else blank_pages).append(index)
    if ocr_pages:
        if not ocr_enabled:
            if not units_by_page:
                raise ParseError("OCR_REQUIRED", "PDF 未包含可提取文本，当前配置未启用 OCR")
        elif len(ocr_pages) > max_ocr_pages:
            raise ParseError("OCR_PAGE_LIMIT", f"需要 OCR 的页面超过 {max_ocr_pages} 页限制")
        else:
            ocr_units, confidence_values = _ocr_pdf_pages(data, ocr_pages)
            units_by_page.update(ocr_units)
    if not units_by_page:
        raise ParseError("NO_EXTRACTABLE_TEXT", "PDF 未提取到可检索文字")
    quality = "ocr-reviewed" if confidence_values and min(confidence_values) >= 0.70 else "ocr-low-confidence" if confidence_values else "native"
    parser_name = "pypdf+rapidocr" if confidence_values else "pypdf"
    metadata = {
        "pageCount": len(reader.pages), "ocrPageCount": len(confidence_values),
        "blankPageCount": len(blank_pages),
        "ocrMeanConfidence": round(sum(confidence_values) / len(confidence_values), 4) if confidence_values else None,
    }
    units = [units_by_page[index] for index in sorted(units_by_page)]
    return ParsedDocument("\n\n".join(units), MIME_BY_EXTENSION[".pdf"], parser_name, locator="PDF 正文", extraction_quality=quality, metadata=metadata)


def _ocr_pdf_pages(data: bytes, page_numbers: list[int]) -> tuple[dict[int, str], list[float]]:
    try:
        import numpy as np
        import pypdfium2 as pdfium
        from rapidocr import RapidOCR
    except ImportError as error:
        raise ParseError("OCR_UNAVAILABLE", "本机 OCR 组件不可用，请安装 rapidocr、onnxruntime 和 pypdfium2") from error

    global _OCR_ENGINE
    with _OCR_LOCK:
        if _OCR_ENGINE is None:
            try:
                _OCR_ENGINE = RapidOCR()
            except Exception as error:
                raise ParseError("OCR_UNAVAILABLE", "本机 OCR 模型加载失败") from error
    pdf = pdfium.PdfDocument(data)
    units: dict[int, str] = {}
    confidences: list[float] = []
    try:
        for page_number in page_numbers:
            page = pdf[page_number - 1]
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil().convert("RGB")
            try:
                with _OCR_LOCK:
                    result = _OCR_ENGINE(np.asarray(image))
                texts = [str(item).strip() for item in (getattr(result, "txts", None) or []) if str(item).strip()]
                scores = [float(item) for item in (getattr(result, "scores", None) or [])]
                if texts:
                    units[page_number] = f"[[LOCATOR:第 {page_number} 页 · OCR]]\n" + "\n".join(texts)
                    confidences.append(sum(scores) / len(scores) if scores else 0.0)
            finally:
                image.close()
                bitmap.close()
                page.close()
    finally:
        pdf.close()
    if len(units) != len(page_numbers):
        raise ParseError("OCR_PARTIAL", "部分扫描页未识别到文字，文件未写入知识库")
    return units, confidences


def _parse_docx(data: bytes) -> ParsedDocument:
    document = Document(io.BytesIO(data))
    units: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            style = (paragraph.style.name or "").lower()
            locator = f"标题：{text[:60]}" if "heading" in style or "标题" in style else f"段落 {index}"
            units.append(f"[[LOCATOR:{locator}]]\n{text}")
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        headers = rows[0]
        for row_index, row in enumerate(rows[1:] or rows, start=1):
            values = "；".join(f"{headers[i] if i < len(headers) else f'列{i + 1}'}：{value}" for i, value in enumerate(row) if value)
            if values:
                units.append(f"[[LOCATOR:表格 {table_index} · 第 {row_index} 行]]\n{values}")
    return ParsedDocument("\n\n".join(units), MIME_BY_EXTENSION[".docx"], "python-docx", locator="DOCX 正文")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip().replace("\n", " ")


def _parse_xlsx(data: bytes, *, max_sheets: int, max_rows: int, max_columns: int) -> ParsedDocument:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) > max_sheets:
            raise ParseError("XLSX_SHEET_LIMIT", f"Excel 工作表超过 {max_sheets} 个限制")
        units: list[str] = []
        total_rows = 0
        for worksheet in workbook.worksheets:
            if worksheet.max_row > max_rows:
                raise ParseError("XLSX_ROW_LIMIT", f"工作表“{worksheet.title}”超过 {max_rows} 行限制")
            if worksheet.max_column > max_columns:
                raise ParseError("XLSX_COLUMN_LIMIT", f"工作表“{worksheet.title}”超过 {max_columns} 列限制")
            rows = worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, max_col=worksheet.max_column, values_only=True)
            first = next(rows, None)
            if first is None:
                continue
            headers = [_cell_text(value) or f"列 {index + 1}" for index, value in enumerate(first)]
            for row_number, row in enumerate(rows, start=2):
                values = [_cell_text(value) for value in row]
                pairs = [f"{headers[index] if index < len(headers) else f'列 {index + 1}'}：{value}" for index, value in enumerate(values) if value]
                if not pairs:
                    continue
                end_column = get_column_letter(max(1, len(values)))
                units.append(f"[[LOCATOR:工作表 {worksheet.title} · A{row_number}:{end_column}{row_number}]]\n" + "；".join(pairs))
                total_rows += 1
        if not units:
            raise ParseError("NO_EXTRACTABLE_TEXT", "Excel 没有可检索的数据行")
        return ParsedDocument(
            "\n\n".join(units), MIME_BY_EXTENSION[".xlsx"], "openpyxl", locator="Excel 工作簿",
            metadata={"sheetCount": len(workbook.worksheets), "dataRowCount": total_rows},
        )
    finally:
        workbook.close()


def _parse_pptx(data: bytes, *, max_slides: int) -> ParsedDocument:
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > max_slides:
        raise ParseError("PPTX_SLIDE_LIMIT", f"PPT 超过 {max_slides} 页限制")
    units: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text and text != title:
                    slide_parts.append(text)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in shape.table.rows]
                if rows:
                    headers = rows[0]
                    for row_index, row in enumerate(rows[1:] or rows, start=1):
                        values = "；".join(f"{headers[index] if index < len(headers) and headers[index] else f'列 {index + 1}'}：{value}" for index, value in enumerate(row) if value)
                        if values:
                            slide_parts.append(f"表格第 {row_index} 行：{values}")
        try:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except (AttributeError, ValueError):
            notes = ""
        if notes:
            slide_parts.append(f"演讲者备注：{notes}")
        if title or slide_parts:
            heading = f"标题：{title}" if title else ""
            body = "\n".join(part for part in [heading, *slide_parts] if part)
            units.append(f"[[LOCATOR:第 {slide_number} 页]]\n{body}")
    if not units:
        raise ParseError("NO_EXTRACTABLE_TEXT", "PPT 没有可检索文字；纯图片幻灯片暂不自动 OCR")
    return ParsedDocument(
        "\n\n".join(units), MIME_BY_EXTENSION[".pptx"], "python-pptx", locator="PPT 幻灯片",
        metadata={"slideCount": len(presentation.slides)},
    )


def _parse_csv(data: bytes) -> ParsedDocument:
    rows = list(csv.reader(io.StringIO(_decode_text(data))))
    if not rows:
        raise ParseError("EMPTY_FILE", "CSV 文件为空")
    headers = [item.strip() or f"列 {index + 1}" for index, item in enumerate(rows[0])]
    units = []
    for index, row in enumerate(rows[1:], start=2):
        values = "；".join(f"{headers[column] if column < len(headers) else f'列 {column + 1}'}：{value.strip()}" for column, value in enumerate(row) if value.strip())
        if values:
            units.append(f"[[LOCATOR:第 {index} 行]]\n{values}")
    if not units:
        raise ParseError("NO_EXTRACTABLE_TEXT", "CSV 没有可检索的数据行")
    return ParsedDocument("\n\n".join(units), MIME_BY_EXTENSION[".csv"], "csv", locator="CSV 数据")


def _parse_html(data: bytes) -> ParsedDocument:
    parser = _HtmlText()
    parser.feed(_decode_text(data))
    return ParsedDocument("\n".join(parser.items), MIME_BY_EXTENSION[".html"], "html.parser", locator="HTML 正文")
