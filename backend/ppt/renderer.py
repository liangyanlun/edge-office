from __future__ import annotations

import os
import base64
import json
import threading
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

RENDER_REVISION = 2
_OFFICE_LOCK = threading.Lock()


def hidden_process_options():
    if os.name != "nt":
        return {}
    startup = subprocess.STARTUPINFO()
    startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    startup.wShowWindow = 0
    return {"startupinfo": startup, "creationflags": subprocess.CREATE_NO_WINDOW}


def render_presentation(pptx_path: Path, output_dir: Path, allow_external: bool = True) -> dict[str, Any]:
    """Render previews with Office in an isolated process, then degrade safely."""
    output_dir.mkdir(parents=True, exist_ok=True)
    for old in output_dir.glob("slide-*.png"):
        old.unlink(missing_ok=True)
    deck = Presentation(str(pptx_path))
    count = len(deck.slides)
    mode = None
    def complete():
        try:
            for index in range(1, count + 1):
                with Image.open(output_dir / f"slide-{index}.png") as image:
                    image.verify()
            return count > 0
        except (OSError, ValueError):
            return False
    if allow_external:
        with _OFFICE_LOCK:
            if os.name == "nt" and _render_with_powerpoint_worker(pptx_path, output_dir) and complete():
                mode = "powerpoint-com-worker"
            else:
                for old in output_dir.glob("slide-*.png"):
                    old.unlink(missing_ok=True)
                if _render_with_libreoffice(pptx_path, output_dir) and complete():
                    mode = "libreoffice"
    if mode is None:
        for old in output_dir.glob("slide-*.png"):
            old.unlink(missing_ok=True)
        _render_approximate(pptx_path, output_dir)
        mode = "safe-layout-preview"
    result = {"revision": RENDER_REVISION, "mode": mode, "approximate": mode == "safe-layout-preview", "count": count,
              "slideSize": {"width": int(deck.slide_width), "height": int(deck.slide_height)}}
    if result["approximate"]:
        result["warning"] = "当前为结构预览；中文与图片可核对，但母版、图表、SmartArt 和复杂样式请在 PowerPoint 中检查。"
    (output_dir / "render.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return result


def _render_with_powerpoint_worker(pptx_path: Path, output_dir: Path) -> bool:
    executable = shutil.which("powershell") or shutil.which("pwsh")
    if not executable:
        return False
    script = r"""
$ErrorActionPreference='Stop'
$inputPath=$env:EDGE_PPT_INPUT; $outputPath=$env:EDGE_PPT_OUTPUT
$ppt=$null; $deck=$null
try {
  $ppt=New-Object -ComObject PowerPoint.Application
  $ppt.AutomationSecurity=3
  $deck=$ppt.Presentations.Open($inputPath,-1,0,0)
  $height=[int][Math]::Round(1280*$deck.PageSetup.SlideHeight/$deck.PageSetup.SlideWidth)
  foreach($slide in $deck.Slides) { $slide.Export((Join-Path $outputPath ("slide-{0}.png" -f $slide.SlideIndex)),"PNG",1280,$height) }
} finally { if($deck){$deck.Close()}; if($ppt){$ppt.Quit()} }
"""
    try:
        result = subprocess.run(
            [executable, "-NoProfile", "-NonInteractive", "-WindowStyle", "Hidden", "-EncodedCommand", base64.b64encode(script.encode("utf-16-le")).decode("ascii")],
            env={**os.environ, "EDGE_PPT_INPUT": str(pptx_path.resolve()), "EDGE_PPT_OUTPUT": str(output_dir.resolve())},
            capture_output=True, timeout=60, check=False, **hidden_process_options(),
        )
        return result.returncode == 0 and bool(list(output_dir.glob("slide-*.png")))
    except (OSError, subprocess.SubprocessError):
        return False


def _render_with_libreoffice(pptx_path: Path, output_dir: Path) -> bool:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable and os.name == "nt":
        for base in (os.getenv("ProgramFiles", ""), os.getenv("ProgramFiles(x86)", "")):
            candidate = Path(base) / "LibreOffice/program/soffice.exe"
            if candidate.is_file():
                executable = str(candidate)
                break
    if not executable:
        return False
    try:
        import pypdfium2 as pdfium
        with tempfile.TemporaryDirectory(prefix="edge-office-ppt-") as temp:
            result = subprocess.run(
                [executable, f"-env:UserInstallation={(Path(temp) / 'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", temp, str(pptx_path.resolve())],
                capture_output=True, timeout=60, check=False, **hidden_process_options(),
            )
            pdf_path = Path(temp) / f"{pptx_path.stem}.pdf"
            if result.returncode != 0 or not pdf_path.is_file():
                return False
            document = pdfium.PdfDocument(str(pdf_path))
            try:
                for index in range(len(document)):
                    page = document[index]
                    bitmap = page.render(scale=1.5)
                    bitmap.to_pil().convert("RGB").save(output_dir / f"slide-{index + 1}.png")
                    bitmap.close()
                    page.close()
            finally:
                document.close()
        return bool(list(output_dir.glob("slide-*.png")))
    except Exception:
        return False


def _render_approximate(pptx_path: Path, output_dir: Path) -> int:
    presentation = Presentation(str(pptx_path))
    width = 1280
    height = round(width * presentation.slide_height / presentation.slide_width)
    point_scale = width / (presentation.slide_width / 914400) / 72
    for index, slide in enumerate(presentation.slides, start=1):
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        for shape in slide.shapes:
            left = round(shape.left / presentation.slide_width * width)
            top = round(shape.top / presentation.slide_height * height)
            right = round((shape.left + shape.width) / presentation.slide_width * width)
            bottom = round((shape.top + shape.height) / presentation.slide_height * height)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    with Image.open(__import__("io").BytesIO(shape.image.blob)) as source:
                        source.thumbnail((max(1, right - left), max(1, bottom - top)))
                        image.paste(source.convert("RGB"), (left, top))
                except Exception:
                    draw.rectangle((left, top, right, bottom), outline="#98a5b5", width=2)
            elif getattr(shape, "has_text_frame", False):
                layer = Image.new("RGBA", (max(1, right-left), max(1, bottom-top)))
                text_draw = ImageDraw.Draw(layer)
                y = 4
                for paragraph in shape.text_frame.paragraphs:
                    size = paragraph.font.size or next((r.font.size for r in paragraph.runs if r.font.size), None)
                    font = preview_font(round((size.pt if size else 20) * point_scale))
                    line = ""
                    for char in paragraph.text[:20000] + "\n":
                        if char == "\n" or (line and text_draw.textlength(line + char, font=font) > max(1, right-left-8)):
                            text_draw.text((4, y), line, fill="#1f2c3d", font=font)
                            y += round(font.size * 1.25)
                            line = ""
                        if char != "\n":
                            line += char
                        if y > bottom-top:
                            break
                image.paste(layer, (left, top), layer)
            else:
                draw.rectangle((left, top, right, bottom), outline="#d9dfe7", width=1)
        image.save(output_dir / f"slide-{index}.png")
    return len(presentation.slides)


def preview_font(size):
    # A model/PPT-supplied font name never becomes a local filesystem path.
    for path in (Path(os.getenv("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc",
                 Path("C:/Windows/Fonts/simsun.ttc"),
                 Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
                 Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")):
        if path.is_file():
            return ImageFont.truetype(str(path), size=max(8, min(size, 160)))
    return ImageFont.load_default(size=max(8, min(size, 160)))
