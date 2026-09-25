from __future__ import annotations

import hashlib
import html
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


ALLOWED_TEXT_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
}


def build_snapshot(pptx_path: Path, base_version: str) -> dict[str, Any]:
    presentation = Presentation(str(pptx_path))
    slides = [_slide_snapshot(slide, index + 1, base_version) for index, slide in enumerate(presentation.slides)]
    return {
        "schema": "edge-office-ppt-snapshot/1.0",
        "baseVersion": base_version,
        "slideSize": {"width": presentation.slide_width, "height": presentation.slide_height},
        "slides": slides,
    }


def snapshot_for_slide(snapshot: dict[str, Any], slide_id: str) -> dict[str, Any]:
    for slide in snapshot.get("slides", []):
        if slide.get("slideId") == slide_id:
            return slide
    raise KeyError(slide_id)


def write_snapshot(snapshot: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")


def load_snapshot(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def slide_structure_digest(slide: dict[str, Any]) -> str:
    canonical = json.dumps(
        {
            "slideId": slide.get("slideId"),
            "elements": [
                {
                    "kind": item.get("kind"),
                    "id": item.get("id"),
                    "paragraphCount": len(item.get("paragraphs", [])),
                    "paragraphs": [p.get("text", "") for p in item.get("paragraphs", [])],
                    "imageSha256": item.get("imageSha256"),
                    "bounds": item.get("bounds"),
                }
                for item in slide.get("elements", [])
            ],
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _slide_snapshot(slide: Any, slide_number: int, base_version: str) -> dict[str, Any]:
    elements: list[dict[str, Any]] = []
    html_parts = [f'<section data-slide-id="slide-{slide_number}" data-base-version="{html.escape(base_version)}">']
    for shape in slide.shapes:
        shape_id = int(shape.shape_id)
        bounds = {"left": int(shape.left), "top": int(shape.top), "width": int(shape.width), "height": int(shape.height)}
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            elements.append({"kind": "unsupported", "id": shape_id, "reason": "group", "bounds": bounds})
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image_bytes = shape.image.blob
            digest = hashlib.sha256(image_bytes).hexdigest()
            elements.append({
                "kind": "image", "id": shape_id, "imageId": shape_id, "name": shape.name,
                "imageSha256": digest, "contentType": shape.image.content_type, "bounds": bounds,
            })
            html_parts.append(f'<img id="{shape_id}" data-sha256="{digest}" alt="{html.escape(shape.name)}" />')
            continue
        if getattr(shape, "has_text_frame", False):
            placeholder_type = None
            editable = True
            if getattr(shape, "is_placeholder", False):
                placeholder_type = int(shape.placeholder_format.type)
                editable = shape.placeholder_format.type in ALLOWED_TEXT_PLACEHOLDERS
            paragraphs = [
                {
                    "id": index,
                    "paragraphId": index,
                    "text": paragraph.text,
                    "runs": [run.text for run in paragraph.runs],
                }
                for index, paragraph in enumerate(shape.text_frame.paragraphs)
            ]
            elements.append({
                "kind": "text", "id": shape_id, "divId": shape_id, "name": shape.name,
                "paragraphs": paragraphs, "editable": editable, "placeholderType": placeholder_type, "bounds": bounds,
            })
            paragraph_html = "".join(
                f'<p id="{item["id"]}">{html.escape(item["text"])}</p>' for item in paragraphs
            )
            html_parts.append(f'<div id="{shape_id}" data-editable="{str(editable).lower()}">{paragraph_html}</div>')
        elif shape.shape_type not in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE}:
            elements.append({"kind": "unsupported", "id": shape_id, "reason": str(shape.shape_type), "bounds": bounds})
    html_parts.append("</section>")
    result = {
        "slideId": f"slide-{slide_number}",
        "slideNumber": slide_number,
        "baseVersion": base_version,
        "elements": elements,
        "html": "".join(html_parts),
    }
    result["structureDigest"] = slide_structure_digest(result)
    return result
