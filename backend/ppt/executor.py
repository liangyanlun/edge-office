from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .contracts import PptContractError, PptPatch


def execute_patch(base_path: Path, candidate_path: Path, patch: PptPatch, operations: list[dict[str, Any]]) -> None:
    candidate_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(base_path, candidate_path)
    try:
        candidate_path.chmod(0o666)
    except OSError:
        pass
    presentation = Presentation(str(candidate_path))
    slide_index = _slide_index(patch.slide_id)
    if slide_index >= len(presentation.slides):
        raise PptContractError("PPT_SLIDE_NOT_FOUND", "目标幻灯片不存在")
    slide = presentation.slides[slide_index]
    shapes = {int(shape.shape_id): shape for shape in slide.shapes}
    paragraph_maps: dict[int, dict[int, Any]] = {}
    for shape_id, shape in shapes.items():
        if getattr(shape, "has_text_frame", False):
            paragraph_maps[shape_id] = {index: paragraph for index, paragraph in enumerate(shape.text_frame.paragraphs)}
    for operation in operations:
        name = operation["operation"]
        if name in {"replace_paragraph", "clone_paragraph", "del_paragraph"}:
            div_id = int(operation["div_id"])
            paragraph_id = int(operation["paragraph_id"])
            mapping = paragraph_maps[div_id]
            paragraph = mapping.get(paragraph_id)
            if paragraph is None:
                raise PptContractError("PPT_PARAGRAPH_NOT_FOUND", f"未找到段落 {paragraph_id}")
            if name == "replace_paragraph":
                _replace_paragraph_text(paragraph, str(operation["text"]))
            elif name == "clone_paragraph":
                clone = deepcopy(paragraph._p)
                paragraph._p.addnext(clone)
                new_id = max(mapping, default=-1) + 1
                mapping[new_id] = _paragraph_proxy(shapes[div_id], clone)
            else:
                paragraph._p.getparent().remove(paragraph._p)
                del mapping[paragraph_id]
        elif name == "del_image":
            image_id = int(operation["image_id"])
            shape = shapes[image_id]
            shape._element.getparent().remove(shape._element)
            del shapes[image_id]
        elif name == "replace_image":
            old = shapes[int(operation["image_id"])]
            if old.shape_type != MSO_SHAPE_TYPE.PICTURE:
                raise PptContractError("PPT_IMAGE_TARGET_NOT_FOUND", "目标对象不是图片")
            new = slide.shapes.add_picture(str(operation["assetPath"]), old.left, old.top, old.width, old.height)
            new.rotation = old.rotation
            for attr in ("crop_left", "crop_right", "crop_top", "crop_bottom"):
                try:
                    setattr(new, attr, getattr(old, attr))
                except (AttributeError, ValueError):
                    pass
            parent = old._element.getparent()
            index = parent.index(old._element)
            parent.remove(new._element)
            parent.insert(index, new._element)
            parent.remove(old._element)
            shapes[int(operation["image_id"])] = new
    presentation.save(str(candidate_path))


def _replace_paragraph_text(paragraph: Any, text: str) -> None:
    runs = list(paragraph.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run().text = text


def _paragraph_proxy(shape: Any, paragraph_xml: Any) -> Any:
    for paragraph in shape.text_frame.paragraphs:
        if paragraph._p is paragraph_xml:
            return paragraph
    raise RuntimeError("克隆段落未能重新绑定")


def _slide_index(slide_id: str) -> int:
    if not slide_id.startswith("slide-"):
        raise PptContractError("PPT_SLIDE_ID_INVALID", "幻灯片 ID 格式不正确")
    try:
        value = int(slide_id.removeprefix("slide-")) - 1
    except ValueError as exc:
        raise PptContractError("PPT_SLIDE_ID_INVALID", "幻灯片 ID 格式不正确") from exc
    if value < 0:
        raise PptContractError("PPT_SLIDE_ID_INVALID", "幻灯片 ID 格式不正确")
    return value
