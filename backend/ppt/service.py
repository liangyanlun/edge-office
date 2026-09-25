from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation

from ..database import Database
from ..inference import InferenceAdapter
from ..parsers import ParseError, validate_ooxml_archive
from .contracts import PptContractError, PptOperation, PptPatch
from .executor import execute_patch
from .planner import PptPlanner
from .policy import validate_patch
from .renderer import render_presentation
from .snapshot import build_snapshot, load_snapshot, snapshot_for_slide, write_snapshot
from .storage import PptStorage
from .verifier import verify_candidate


class PptServiceError(RuntimeError):
    def __init__(self, code: str, message: str, status: int = 400):
        super().__init__(message)
        self.code = code
        self.status = status


class PptService:
    def __init__(self, database: Database, data_dir: Path, inference: InferenceAdapter, max_upload_bytes: int, max_slides: int = 300, allow_external_renderer: bool = True):
        self.storage = PptStorage(database, data_dir)
        self.planner = PptPlanner(inference)
        self.max_upload_bytes = max_upload_bytes
        self.max_slides = max_slides
        self.allow_external_renderer = allow_external_renderer

    def runtime(self) -> dict[str, Any]:
        return {
            "capability": self.planner.capability(),
            "operations": ["replace_paragraph", "clone_paragraph", "del_paragraph", "replace_image", "del_image"],
            "limits": {"maxOperations": 32, "maxUploadBytes": self.max_upload_bytes},
        }

    def import_presentation(self, name: str, data: bytes) -> dict[str, Any]:
        if not name.lower().endswith(".pptx"):
            raise PptServiceError("PPTX_REQUIRED", "PPT 编辑仅接受 .pptx 文件")
        if not data:
            raise PptServiceError("PPT_FILE_EMPTY", "PPTX 文件为空")
        if len(data) > self.max_upload_bytes:
            raise PptServiceError("PPT_FILE_TOO_LARGE", "PPTX 文件超过上传限制", 413)
        try:
            validate_ooxml_archive(data)
        except ParseError as exc:
            raise PptServiceError(exc.code, str(exc)) from exc
        prepared = self.storage.prepare_import(name, data)
        try:
            presentation = Presentation(str(prepared["versionPath"]))
            if not presentation.slides:
                raise PptServiceError("PPT_EMPTY", "PPTX 不包含幻灯片")
            if len(presentation.slides) > self.max_slides:
                raise PptServiceError("PPT_SLIDE_LIMIT", f"PPTX 超过 {self.max_slides} 页限制")
            snapshot = build_snapshot(prepared["versionPath"], prepared["versionId"])
            snapshot_path = prepared["versionDir"] / "snapshot.json"
            write_snapshot(snapshot, snapshot_path)
            render = render_presentation(prepared["versionPath"], prepared["versionDir"] / "slides", self.allow_external_renderer)
            item = self.storage.commit_import(prepared, snapshot_path)
            return {**item, "slides": self._slides(snapshot), "render": render, "modelCapability": self.planner.capability()}
        except PptServiceError:
            self.storage.cleanup_prepared(prepared)
            raise
        except Exception as exc:
            self.storage.cleanup_prepared(prepared)
            raise PptServiceError("PPT_IMPORT_FAILED", f"PPTX 解析失败：{exc}") from exc

    def list_presentations(self) -> list[dict[str, Any]]:
        return self.storage.list_presentations()

    def get_presentation(self, presentation_id: str) -> dict[str, Any]:
        item = self.storage.get_presentation(presentation_id)
        if not item:
            raise PptServiceError("PPT_NOT_FOUND", "演示文稿不存在", 404)
        item["modelCapability"] = self.planner.capability()
        return item

    def slides(self, presentation_id: str) -> dict[str, Any]:
        version = self._version(presentation_id)
        snapshot = load_snapshot(Path(version["snapshotPath"]))
        return {"presentationId": presentation_id, "versionId": version["id"], "items": self._slides(snapshot)}

    def slide_snapshot(self, presentation_id: str, slide_id: str) -> dict[str, Any]:
        version = self._version(presentation_id)
        try:
            slide = snapshot_for_slide(load_snapshot(Path(version["snapshotPath"])), slide_id)
        except KeyError as exc:
            raise PptServiceError("PPT_SLIDE_NOT_FOUND", "幻灯片不存在", 404) from exc
        return {"presentationId": presentation_id, "versionId": version["id"], "item": slide}

    def upload_asset(self, presentation_id: str, name: str, mime_type: str, data: bytes) -> dict[str, Any]:
        if mime_type not in {"image/png", "image/jpeg", "image/webp"}:
            raise PptServiceError("PPT_IMAGE_TYPE_UNSUPPORTED", "替换图片仅支持 PNG、JPEG 或 WebP")
        if not data or len(data) > 12 * 1024 * 1024:
            raise PptServiceError("PPT_IMAGE_SIZE_INVALID", "替换图片不能为空且不能超过 12MB")
        try:
            with Image.open(BytesIO(data)) as image:
                image.verify()
            with Image.open(BytesIO(data)) as image:
                if image.width * image.height > 40_000_000:
                    raise PptServiceError("PPT_IMAGE_PIXELS_EXCEEDED", "替换图片像素总量超过限制")
        except PptServiceError:
            raise
        except Exception as exc:
            raise PptServiceError("PPT_IMAGE_INVALID", "替换图片内容损坏或格式不正确") from exc
        try:
            return self.storage.save_asset(presentation_id, name, mime_type, data)
        except KeyError as exc:
            raise PptServiceError("PPT_NOT_FOUND", "演示文稿不存在", 404) from exc

    def create_patch(self, presentation_id: str, payload: dict[str, Any], *, api_sequence: str | None = None) -> dict[str, Any]:
        presentation = self.get_presentation(presentation_id)
        base_version = str(payload.get("baseVersion") or presentation["currentVersionId"])
        if base_version != presentation["currentVersionId"]:
            raise PptServiceError("PPT_BASE_VERSION_STALE", "当前页面基于旧版本，请刷新后重新生成计划", 409)
        slide_id = str(payload.get("slideId") or "")
        instruction = str(payload.get("instruction") or "").strip()
        if not slide_id or not instruction or len(instruction) > 4000:
            raise PptServiceError("PPT_PATCH_INPUT_INVALID", "请选择页面并输入不超过 4000 字的修改要求")
        version = self._version(presentation_id, base_version)
        snapshot = load_snapshot(Path(version["snapshotPath"]))
        try:
            slide = snapshot_for_slide(snapshot, slide_id)
            assets = self.storage.list_assets(presentation_id)
            patch, raw, capability = self.planner.plan(presentation_id, base_version, slide_id, instruction, slide, assets, api_sequence)
            validate_patch(patch, snapshot, self.storage.get_asset)
            item = self.storage.create_patch(patch, instruction, capability["model"], raw, capability["warning"])
        except KeyError as exc:
            raise PptServiceError("PPT_SLIDE_NOT_FOUND", "幻灯片不存在", 404) from exc
        except PptContractError as exc:
            raise PptServiceError(exc.code, str(exc), 422) from exc
        except RuntimeError as exc:
            raise PptServiceError("PPT_PLANNER_UNAVAILABLE", str(exc), 503) from exc
        return self._public_patch(item)

    def get_patch(self, patch_id: str) -> dict[str, Any]:
        patch = self.storage.get_patch(patch_id)
        if not patch:
            raise PptServiceError("PPT_PATCH_NOT_FOUND", "PPT 修改计划不存在", 404)
        return self._public_patch(patch)

    def simulate(self, patch_id: str) -> dict[str, Any]:
        patch_row = self._patch(patch_id)
        if patch_row["status"] not in {"PLANNED", "FAILED"}:
            if patch_row["status"] == "PREVIEW_READY":
                return self._public_patch(patch_row)
            raise PptServiceError("PPT_PATCH_STATE_INVALID", "当前计划不能生成预览", 409)
        presentation = self.get_presentation(patch_row["presentationId"])
        if presentation["currentVersionId"] != patch_row["baseVersion"]:
            raise PptServiceError("PPT_BASE_VERSION_STALE", "计划对应的基础版本已过期", 409)
        version = self._version(patch_row["presentationId"], patch_row["baseVersion"])
        base_snapshot = load_snapshot(Path(version["snapshotPath"]))
        contract = self._contract(patch_row)
        try:
            resolved = validate_patch(contract, base_snapshot, self.storage.get_asset)
            candidate = self.storage.staged_path(patch_row["presentationId"], patch_id)
            execute_patch(Path(version["path"]), candidate, contract, resolved)
            staged_snapshot = candidate.parent / "snapshot.json"
            verification = verify_candidate(candidate, base_snapshot, contract, resolved, staged_snapshot)
            render = render_presentation(candidate, candidate.parent / "preview", self.allow_external_renderer)
            verification["render"] = render
            self.storage.update_patch(patch_id, "PREVIEW_READY", staged_path=candidate, verification=verification, event="ppt_preview_ready", event_data={"verification": verification})
            self.storage.update_patch(patch_id, "PREVIEW_READY", event="ppt_verification", event_data={"ok": True, "candidateSha256": verification["candidateSha256"]})
            return self._public_patch(self._patch(patch_id))
        except Exception as exc:
            self.storage.delete_staged(patch_row)
            code = exc.code if isinstance(exc, PptContractError) else "PPT_EXECUTION_FAILED"
            self.storage.update_patch(patch_id, "FAILED", verification={"ok": False, "code": code}, event="ppt_verification", event_data={"ok": False, "code": code})
            raise PptServiceError(code, str(exc), 422) from exc

    def confirm(self, patch_id: str) -> dict[str, Any]:
        patch = self._patch(patch_id)
        if patch["status"] != "PREVIEW_READY" or not patch["stagedPath"]:
            raise PptServiceError("PPT_PATCH_NOT_PREVIEWED", "只能发布已生成并验证过的预览", 409)
        contract = self._contract(patch)
        if contract.canonical_json() != patch["canonicalJson"]:
            raise PptServiceError("PPT_PATCH_TAMPERED", "PPT 修改计划规范化内容校验失败", 409)
        if hashlib.sha256(patch["canonicalJson"].encode("utf-8")).hexdigest() != patch["patchHash"]:
            raise PptServiceError("PPT_PATCH_TAMPERED", "PPT 修改计划哈希校验失败", 409)
        candidate = Path(patch["stagedPath"])
        digest = hashlib.sha256(candidate.read_bytes()).hexdigest()
        if digest != patch["verification"].get("candidateSha256"):
            raise PptServiceError("PPT_STAGED_FILE_TAMPERED", "暂存文件与已确认预览不一致", 409)
        version_id = self.storage.new_version_id(patch["presentationId"], digest)
        final_snapshot = candidate.parent / "publish-snapshot.json"
        write_snapshot(build_snapshot(candidate, version_id), final_snapshot)
        try:
            version = self.storage.publish_patch(patch, version_id, final_snapshot, candidate.parent / "preview", digest)
        except RuntimeError as exc:
            if str(exc) == "PPT_BASE_VERSION_STALE":
                raise PptServiceError("PPT_BASE_VERSION_STALE", "基础版本已变化，不能发布旧预览", 409) from exc
            raise
        self.storage.update_patch(patch_id, "PUBLISHED", published_version_id=version["id"], event="ppt_complete", event_data={"versionId": version["id"]})
        return {"patch": self._public_patch(self._patch(patch_id)), "version": version}

    def cancel(self, patch_id: str) -> dict[str, Any]:
        patch = self._patch(patch_id)
        if patch["status"] in {"PUBLISHED", "CANCELLED"}:
            raise PptServiceError("PPT_PATCH_STATE_INVALID", "当前计划不能取消", 409)
        self.storage.delete_staged(patch)
        self.storage.update_patch(patch_id, "CANCELLED", event="ppt_cancelled", event_data={})
        return self._public_patch(self._patch(patch_id))

    def restore(self, presentation_id: str, version_id: str) -> dict[str, Any]:
        try:
            restored = self.storage.restore_version(presentation_id, version_id)
        except KeyError as exc:
            raise PptServiceError("PPT_VERSION_NOT_FOUND", "版本不存在", 404) from exc
        version = self._version(presentation_id, restored["id"])
        write_snapshot(build_snapshot(Path(version["path"]), restored["id"]), Path(version["snapshotPath"]))
        return restored

    def version_file(self, presentation_id: str, version_id: str) -> Path:
        return Path(self._version(presentation_id, version_id)["path"])

    def preview_file(self, presentation_id: str, version_id: str, slide_id: str) -> Path:
        version = self._version(presentation_id, version_id)
        path = Path(version["slidesDir"]) / f"{slide_id}.png"
        if not path.is_file():
            raise PptServiceError("PPT_PREVIEW_NOT_FOUND", "页面预览不存在", 404)
        return path

    def patch_preview_file(self, patch_id: str, slide_id: str) -> Path:
        patch = self._patch(patch_id)
        if not patch["stagedPath"]:
            raise PptServiceError("PPT_PREVIEW_NOT_FOUND", "修改后预览尚未生成", 404)
        path = Path(patch["stagedPath"]).parent / "preview" / f"{slide_id}.png"
        if not path.is_file():
            raise PptServiceError("PPT_PREVIEW_NOT_FOUND", "修改后预览不存在", 404)
        return path

    def events(self, patch_id: str) -> list[dict[str, Any]]:
        return self._patch(patch_id)["events"]

    def _version(self, presentation_id: str, version_id: str | None = None) -> dict[str, Any]:
        version = self.storage.get_version(presentation_id, version_id)
        if not version:
            raise PptServiceError("PPT_VERSION_NOT_FOUND", "PPT 版本不存在", 404)
        return version

    def _patch(self, patch_id: str) -> dict[str, Any]:
        patch = self.storage.get_patch(patch_id)
        if not patch:
            raise PptServiceError("PPT_PATCH_NOT_FOUND", "PPT 修改计划不存在", 404)
        return patch

    @staticmethod
    def _contract(patch: dict[str, Any]) -> PptPatch:
        operations = []
        keys = {
            "replace_paragraph": ("div_id", "paragraph_id", "text"), "clone_paragraph": ("div_id", "paragraph_id"),
            "del_paragraph": ("div_id", "paragraph_id"), "replace_image": ("image_id", "image_path"), "del_image": ("image_id",),
        }
        for item in patch["operations"]:
            name = item["operation"]
            operations.append(PptOperation(name, tuple(item[key] for key in keys[name])))
        contract = PptPatch.create(patch["presentationId"], patch["baseVersion"], patch["slideId"], operations)
        if contract.patch_hash != patch["patchHash"]:
            raise PptServiceError("PPT_PATCH_TAMPERED", "PPT 修改计划哈希校验失败", 409)
        return contract

    @staticmethod
    def _slides(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
        return [{"slideId": item["slideId"], "slideNumber": item["slideNumber"], "elementCount": len(item["elements"])} for item in snapshot["slides"]]

    @staticmethod
    def _public_patch(patch: dict[str, Any]) -> dict[str, Any]:
        safe = {key: value for key, value in patch.items() if key not in {"stagedPath", "canonicalJson"}}
        safe["actionPlan"] = {"tool": "ppt_apply_patch", "arguments": {"patch_id": patch["id"]}, "confirmed": patch["status"] == "PUBLISHED"}
        safe["previewReady"] = patch["status"] in {"PREVIEW_READY", "PUBLISHED"}
        return safe
